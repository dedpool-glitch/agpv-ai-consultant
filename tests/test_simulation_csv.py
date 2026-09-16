import csv
from io import StringIO
import unittest
from io import BytesIO

from openpyxl import load_workbook

from services.simulation_csv import build_daily_yield_csv, build_run_summary_csv
from services.simulation_workbook import build_simulation_workbook


def rows(data):
    return list(csv.DictReader(StringIO(data.decode("utf-8-sig"))))


class SimulationCsvTests(unittest.TestCase):
    def test_summary_preserves_raw_values_inputs_and_origins_across_runs(self):
        runs = [
            {
                "label": "Initial",
                "location_context": {"confirmed_address": "Lafayette, Indiana"},
                "input": {"lat": 40.4191229, "lon": -86.8919011,
                          "array": {"pitch": 8.0}},
                "output": {"yearly_yield": 989.6134,
                           "monthly_yield": [10.125] * 12,
                           "yield_unit": "kWh/m"},
                "input_provenance": {"parameters": {"pitch": {
                    "source": "llm_recommended"}}},
            },
            {
                "label": "Requested rerun",
                "input": {"array": {"pitch": 12.0}},
                "output": {"yearly_yield": 1032.61,
                           "monthly_yield": [11.25] * 12,
                           "yield_unit": "kWh/m"},
                "input_provenance": {"parameters": {"pitch": {
                    "source": "user_provided", "user_quote": "Use 12 meters"}}},
            },
        ]
        exported = rows(build_run_summary_csv(runs))
        self.assertEqual([r["run_id"] for r in exported], ["1", "2"])
        self.assertEqual(exported[0]["yearly_yield"], "989.6134")
        self.assertEqual(exported[0]["input_lat"], "40.4191229")
        self.assertEqual(exported[0]["monthly_yield_january"], "10.125")
        self.assertEqual(exported[1]["input_array_pitch"], "12.0")
        self.assertEqual(exported[1]["input_array_pitch_unit"], "m")
        self.assertEqual(exported[1]["input_array_pitch_origin"], "user_provided")
        self.assertEqual(exported[1]["yield_unit"], "kWh/m")

    def test_daily_rows_use_ordered_index_without_inventing_dates(self):
        exported = rows(build_daily_yield_csv([
            {"output": {"daily_yield": [1.125, 2.25], "yield_unit": "kWh/m"}},
            {"output": {"daily_yield": [3.5], "yield_unit": "kWh/m"}},
        ]))
        self.assertEqual([(r["run_id"], r["day_index"], r["daily_yield"])
                          for r in exported],
                         [("1", "1", "1.125"), ("1", "2", "2.25"),
                          ("2", "1", "3.5")])
        self.assertNotIn("date", exported[0])

    def test_workbook_has_typed_values_and_both_analysis_sheets(self):
        run = {
            "input": {"array": {"config": "tracking", "tilt": 25, "pitch": 11}},
            "output": {"yearly_yield": 1554.67,
                       "monthly_yield": [100.125] * 12,
                       "daily_yield": [3.5, 4.25], "yield_unit": "kWh/m"},
        }
        book = load_workbook(BytesIO(build_simulation_workbook([run])))
        self.assertEqual(book.sheetnames, ["Runs", "Daily yield"])
        runs = book["Runs"]
        values = dict(zip((cell.value for cell in runs[1]),
                          (cell.value for cell in runs[2])))
        self.assertEqual(values["yearly_yield"], 1554.67)
        self.assertEqual(values["input_array_pitch"], 11)
        self.assertIs(values["input_array_tilt_used"], False)
        self.assertEqual(values["monthly_yield_january"], 100.125)
        daily = book["Daily yield"]
        self.assertEqual(daily.max_row, 3)
        self.assertEqual(daily["C3"].value, 4.25)


if __name__ == "__main__":
    unittest.main()
