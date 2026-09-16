"""One analysis workbook containing the saved runs and ordered daily yield."""

from io import BytesIO

from services.simulation_csv import daily_yield_rows, run_summary_rows


def build_simulation_workbook(runs):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    for sheet, (headers, rows) in (
        (workbook.active, run_summary_rows(runs)),
        (workbook.create_sheet("Daily yield"), daily_yield_rows(runs)),
    ):
        sheet.title = "Runs" if sheet is workbook.active else "Daily yield"
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header) for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 31
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="23685A")
            cell.font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column_index, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(column_index)].width = (
                18 if header.endswith("_unit") or header.endswith("_origin")
                else min(max(len(header) + 3, 15), 38)
            )

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
